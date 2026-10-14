"""AI Resume ATS Checker
Upload a resume (PDF, DOCX or TXT) -> get an ATS score and improvement tips.
UI: Streamlit | AI: Google Gemini Flash
"""

import json
import os
import re
from io import BytesIO

import streamlit as st
from docx import Document
from google import genai
from google.genai import types
from pypdf import PdfReader

# ----------------------------------------------------------------------------
# Config
# ----------------------------------------------------------------------------
# Change the model name here (or set GEMINI_MODEL in secrets/env) if Google
# releases a newer Flash model. The fallback is tried if the first one fails.
PRIMARY_MODEL = os.getenv("GEMINI_MODEL", "gemini-3.5-flash")
FALLBACK_MODEL = "gemini-2.5-flash"

MAX_RESUME_CHARS = 15000
MAX_JD_CHARS = 6000
MAX_FILE_MB = 5

PROMPT_TEMPLATE = """You are an expert ATS (Applicant Tracking System) analyst and professional resume reviewer.

Analyze the resume below{jd_clause}. Be honest, specific and strict. Do not invent
information that is not in the resume.

Score each category from 0 to 100:
- formatting: clean structure, standard section headings, no tables/columns/graphics issues, consistent dates, readable length
- keywords: relevant industry/role keywords and hard skills{jd_keywords_note}
- experience: quantified achievements, strong action verbs, clear impact, relevance
- education_skills: clear education and skills sections, certifications, tools
- readability: concise, no typos, no fluff, professional tone, good bullet usage

Then give an overall ATS score from 0 to 100 (weighted: formatting 20%, keywords 30%,
experience 30%, education_skills 10%, readability 10%).

Return ONLY valid JSON with exactly this structure (no markdown, no commentary):
{{
  "overall_score": <int>,
  "category_scores": {{
    "formatting": <int>,
    "keywords": <int>,
    "experience": <int>,
    "education_skills": <int>,
    "readability": <int>
  }},
  "summary": "<2-3 sentence overall assessment>",
  "strengths": ["<short point>", "..."],
  "weaknesses": ["<short point>", "..."],
  "missing_keywords": ["<keyword>", "..."],
  "improvements": [
    {{"priority": "High|Medium|Low", "issue": "<what is wrong>", "fix": "<specific action to take>"}}
  ],
  "bullet_rewrites": [
    {{"original": "<weak bullet from the resume>", "improved": "<stronger rewritten bullet>"}}
  ]
}}

Give 3-6 strengths, 3-6 weaknesses, up to 12 missing keywords, 5-8 improvements and
2-4 bullet rewrites (only rewrite bullets that really exist in the resume; keep facts unchanged,
and use placeholders like [X%] if a metric is missing).
{jd_block}
RESUME:
\"\"\"
{resume}
\"\"\"
"""


# ----------------------------------------------------------------------------
# File reading
# ----------------------------------------------------------------------------
def extract_text(uploaded_file) -> str:
    """Extract plain text from an uploaded PDF, DOCX or TXT file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()

    if name.endswith(".pdf"):
        reader = PdfReader(BytesIO(data))
        if reader.is_encrypted:
            try:
                reader.decrypt("")
            except Exception:
                raise ValueError("This PDF is password protected. Please upload an unlocked copy.")
        pages = [(page.extract_text() or "") for page in reader.pages]
        return "\n".join(pages).strip()

    if name.endswith(".docx"):
        doc = Document(BytesIO(data))
        parts = [p.text for p in doc.paragraphs if p.text.strip()]
        for table in doc.tables:
            for row in table.rows:
                cells = [c.text.strip() for c in row.cells if c.text.strip()]
                if cells:
                    parts.append(" | ".join(cells))
        return "\n".join(parts).strip()

    if name.endswith(".txt"):
        return data.decode("utf-8", errors="ignore").strip()

    raise ValueError("Unsupported file type. Please upload a PDF, DOCX or TXT file.")


# ----------------------------------------------------------------------------
# Gemini
# ----------------------------------------------------------------------------
def get_api_key() -> str:
    """Look for the key in Streamlit secrets, then environment variables."""
    try:
        key = st.secrets.get("GEMINI_API_KEY", "")
    except Exception:  # no secrets file present
        key = ""
    return key or os.getenv("GEMINI_API_KEY", "") or os.getenv("GOOGLE_API_KEY", "")


def build_prompt(resume_text: str, job_description: str) -> str:
    resume_text = resume_text[:MAX_RESUME_CHARS]
    job_description = (job_description or "").strip()[:MAX_JD_CHARS]
    if job_description:
        jd_clause = " against the job description provided"
        jd_keywords_note = " (judge mainly against the job description)"
        jd_block = f'\nJOB DESCRIPTION:\n"""\n{job_description}\n"""\n'
    else:
        jd_clause = " for general ATS compatibility"
        jd_keywords_note = ""
        jd_block = ""
    return PROMPT_TEMPLATE.format(
        jd_clause=jd_clause,
        jd_keywords_note=jd_keywords_note,
        jd_block=jd_block,
        resume=resume_text,
    )


def parse_json_response(text: str) -> dict:
    """Parse the model output into a dict, tolerating ```json fences or extra text."""
    text = (text or "").strip()
    text = re.sub(r"^```(?:json)?\s*|\s*```$", "", text, flags=re.IGNORECASE).strip()
    try:
        return json.loads(text)
    except json.JSONDecodeError:
        start, end = text.find("{"), text.rfind("}")
        if start != -1 and end > start:
            return json.loads(text[start : end + 1])
        raise ValueError("The AI response was not valid JSON. Please try again.")


def _clamp(value, default=0) -> int:
    try:
        return max(0, min(100, int(round(float(value)))))
    except (TypeError, ValueError):
        return default


def normalize_result(raw: dict) -> dict:
    """Make sure every field exists and has the right type, so the UI never crashes."""
    cats = raw.get("category_scores") or {}
    category_scores = {
        k: _clamp(cats.get(k))
        for k in ["formatting", "keywords", "experience", "education_skills", "readability"]
    }
    overall = raw.get("overall_score")
    if overall is None and category_scores:
        weights = {"formatting": 0.2, "keywords": 0.3, "experience": 0.3,
                   "education_skills": 0.1, "readability": 0.1}
        overall = sum(category_scores[k] * w for k, w in weights.items())

    def str_list(key):
        value = raw.get(key) or []
        return [str(v) for v in value if str(v).strip()] if isinstance(value, list) else []

    improvements = []
    for item in raw.get("improvements") or []:
        if isinstance(item, dict):
            improvements.append({
                "priority": str(item.get("priority", "Medium")).title(),
                "issue": str(item.get("issue", "")),
                "fix": str(item.get("fix", "")),
            })

    rewrites = []
    for item in raw.get("bullet_rewrites") or []:
        if isinstance(item, dict) and item.get("original") and item.get("improved"):
            rewrites.append({"original": str(item["original"]), "improved": str(item["improved"])})

    return {
        "overall_score": _clamp(overall),
        "category_scores": category_scores,
        "summary": str(raw.get("summary", "")),
        "strengths": str_list("strengths"),
        "weaknesses": str_list("weaknesses"),
        "missing_keywords": str_list("missing_keywords"),
        "improvements": improvements,
        "bullet_rewrites": rewrites,
    }


def analyze_resume(resume_text: str, job_description: str, api_key: str) -> dict:
    client = genai.Client(api_key=api_key)
    prompt = build_prompt(resume_text, job_description)
    config = types.GenerateContentConfig(
        temperature=0.2,
        response_mime_type="application/json",
    )

    last_error = None
    for model in dict.fromkeys([PRIMARY_MODEL, FALLBACK_MODEL]):  # unique, ordered
        try:
            response = client.models.generate_content(model=model, contents=prompt, config=config)
            return normalize_result(parse_json_response(response.text))
        except Exception as exc:  # try the next model, then report
            last_error = exc
    raise RuntimeError(f"Gemini request failed: {last_error}")


# ----------------------------------------------------------------------------
# UI
# ----------------------------------------------------------------------------
def score_label(score: int) -> str:
    if score >= 80:
        return "Excellent"
    if score >= 65:
        return "Good"
    if score >= 50:
        return "Needs work"
    return "Poor"


def render_results(result: dict) -> None:
    score = result["overall_score"]
    st.subheader("Your ATS Score")
    col1, col2 = st.columns([1, 3])
    col1.metric("Overall", f"{score}/100", score_label(score))
    col2.progress(score / 100)
    if result["summary"]:
        col2.write(result["summary"])

    st.markdown("#### Category breakdown")
    names = {
        "formatting": "Formatting",
        "keywords": "Keywords",
        "experience": "Experience",
        "education_skills": "Education & Skills",
        "readability": "Readability",
    }
    cols = st.columns(len(names))
    for col, (key, label) in zip(cols, names.items()):
        col.metric(label, f"{result['category_scores'][key]}")

    left, right = st.columns(2)
    with left:
        st.markdown("#### Strengths")
        for s in result["strengths"] or ["No strengths listed."]:
            st.markdown(f"- {s}")
    with right:
        st.markdown("#### Weaknesses")
        for w in result["weaknesses"] or ["No weaknesses listed."]:
            st.markdown(f"- {w}")

    if result["missing_keywords"]:
        st.markdown("#### Missing keywords")
        st.write(", ".join(f"`{k}`" for k in result["missing_keywords"]))

    st.markdown("#### How to improve your resume")
    icons = {"High": "🔴", "Medium": "🟠", "Low": "🟢"}
    for item in result["improvements"]:
        icon = icons.get(item["priority"], "🟠")
        with st.expander(f"{icon} {item['priority']}: {item['issue']}"):
            st.write(item["fix"])

    if result["bullet_rewrites"]:
        st.markdown("#### Suggested bullet rewrites")
        for pair in result["bullet_rewrites"]:
            st.markdown(f"**Before:** {pair['original']}")
            st.markdown(f"**After:** {pair['improved']}")
            st.divider()

    st.download_button(
        "Download report (JSON)",
        data=json.dumps(result, indent=2),
        file_name="ats_report.json",
        mime="application/json",
    )


def main() -> None:
    st.set_page_config(page_title="AI Resume ATS Checker", page_icon="📄", layout="wide")
    st.title("📄 AI Resume ATS Checker")
    st.caption("Upload your resume and get an ATS score with practical improvement tips.")

    api_key = get_api_key()
    with st.sidebar:
        st.header("Settings")
        if not api_key:
            api_key = st.text_input("Gemini API key", type="password",
                                    help="Get a free key at https://aistudio.google.com/apikey")
        else:
            st.success("API key loaded")
        st.markdown("---")
        st.caption("Your resume is sent to Google's Gemini API for analysis. "
                   "Don't upload anything you are not comfortable sharing.")

    uploaded = st.file_uploader("Upload your resume", type=["pdf", "docx", "txt"])
    job_description = st.text_area(
        "Job description (optional)",
        height=150,
        placeholder="Paste a job description to check how well your resume matches it...",
    )

    if st.button("Analyze resume", type="primary", disabled=uploaded is None):
        if not api_key:
            st.error("Please add your Gemini API key in the sidebar.")
            return
        if uploaded.size > MAX_FILE_MB * 1024 * 1024:
            st.error(f"File is too large. Maximum size is {MAX_FILE_MB} MB.")
            return

        try:
            with st.spinner("Reading your resume..."):
                text = extract_text(uploaded)
        except Exception as exc:
            st.error(f"Could not read the file: {exc}")
            return

        if len(text) < 100:
            st.error("Very little text was found. If your PDF is a scanned image, "
                     "ATS systems can't read it either - export a text-based PDF or upload a DOCX.")
            return

        try:
            with st.spinner("Analyzing with Gemini..."):
                result = analyze_resume(text, job_description, api_key)
        except Exception as exc:
            st.error(str(exc))
            return

        render_results(result)


if __name__ == "__main__":
    main()
